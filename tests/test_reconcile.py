from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook

from argfx.reconcile.conciliacion import Materiality, flow_identities, stock_controls
from argfx.reconcile.finanzas import find, lines, maturity_profile


def workbook(path: Path) -> Path:
    wb = Workbook()
    a11 = wb.active
    a11.title = "A.1.1"
    a11.append(["Indice"])
    a11.append(["III- DEUDA EN SITUACIÓN DE PAGO NORMAL", 1000.0, 1_480_000.0])
    a11.append(["    ORGANISMOS INTERNACIONALES", 300.0, 444_000.0])
    a11.append(["    ADELANTOS TRANSITORIOS BCRA", 0, 0])
    a11.append(["    ADELANTOS TRANSITORIOS BCRA", 50.0, 74_000.0])
    a22 = wb.create_sheet("A.2.2")
    for label, value in [
        ("I - DEUDA BRUTA", 900.0), ("II - DEUDA ELEGIBLE", 20.0), ("1 - Financiamiento", 300.0),
        ("2 - Amortizaciones", 250.0), ("a) Operaciones", 50.0), ("b) Avales", -5.0),
        ("c) Letras", 0.0), ("d) Ajustes", 35.0), ("e) Ajustes", 0.0),
        ("IV - TOTAL VARIACIONES", 80.0), ("V - DEUDA BRUTA", 1000.0),
        ("VI - DEUDA ELEGIBLE", 20.0), ("VII - DEUDA BRUTA", 980.0),
    ]:  # fmt: skip
        a22.append([label, value, value * 1480])
    a31 = wb.create_sheet("A.3.1")
    a31.append([2026, 2027, "Total"])
    a31.append(["Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre", "Subtotal",
                "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Subtotal"])  # fmt: skip
    a31.append([" TÍTULOS PÚBLICOS", *[10.0] * 6, 60.0, *[20.0] * 6, 120.0, 180.0])
    a31.append(["    CAPITAL", *[8.0] * 6, 48.0, *[15.0] * 6, 90.0, 138.0])
    a31.append(["    INTERÉS", *[2.0] * 6, 12.0, *[5.0] * 6, 30.0, 42.0])
    a31.append([" TOTAL", *[10.0] * 6, 60.0, *[20.0] * 6, 120.0, 180.0])
    out = path / "finanzas.xlsx"
    wb.save(out)
    return out


def test_finanzas_reader_and_profile(tmp_path: Path) -> None:
    x = workbook(tmp_path)
    assert find(lines(x, "A.1.1"), "III-").values == (1000.0, 1_480_000.0)
    prof = maturity_profile(x)
    cap = {r.month: r.usd_thousands for r in prof if r.group == "TÍTULOS PÚBLICOS"
           and r.concept == "capital"}  # fmt: skip
    assert len(cap) == 12 and cap["2026-07"] == 8.0 and cap["2027-06"] == 15.0
    assert {r.group for r in prof} == {"TÍTULOS PÚBLICOS", "TOTAL"}


def test_stock_and_flow_controls(tmp_path: Path) -> None:
    x = workbook(tmp_path)
    inst = tmp_path / "instruments.csv"
    with inst.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["instrument_id", "instrument_family", "outstanding_usd_equiv_sigade"])
        w.writerow(["INS-C-0001", "multilateral", 300_000.0])
        w.writerow(["INS-C-0002", "USD", 650_000.0])
        w.writerow(["INS-C-0003", "adelantos transitorios BCRA (intra sector público)", 49_000.0])
    m = Materiality(abs_musd=0.5, rel=0.0)
    stock = {r["partida"][:20]: r for r in stock_controls(lines(x, "A.1.1"), inst, m)}
    assert not stock["Deuda en situación d"]["material"]
    adv = stock["Adelantos transitori"]
    assert adv["diferencia_musd"] == -0.001 and not adv["material"]
    flows = flow_identities(lines(x, "A.2.2"), m)
    assert [r["material"] for r in flows] == [False, False, False, False]


def test_consolidation_splits_public_holdings_by_method_and_law(tmp_path: Path) -> None:
    from datetime import date

    from argfx.contracts.terms import TERM_FILES
    from argfx.reconcile.consolidacion import build_consolidation

    repo = Path(__file__).resolve().parent.parent / "contratos"
    terms = tmp_path / "contratos"
    terms.mkdir()
    for name in TERM_FILES:  # headers only: no term sheets, every bond is bullet
        header = (repo / name).read_text(encoding="utf-8").splitlines()[0]
        (terms / name).write_text(header + "\n", encoding="utf-8")

    def write(name: str, rows: list[list[object]]) -> Path:
        p = tmp_path / name
        with p.open("w", newline="", encoding="utf-8") as fh:
            csv.writer(fh).writerows(rows)
        return p

    comp = write("comp.csv", [
        ["partida", "millones_usd", "porcentaje", "source_id", "referencia"],
        ["agencias_sector_publico", 100, "", "SRC-0001", "x"],
        ["sector_publico_moneda_total", 100, "", "SRC-0001", "x"],
        ["sector_publico_dolares", "", 60, "SRC-0001", "x"],
        ["sector_publico_legislacion_extranjera", "", 10, "SRC-0001", "x"],
    ])  # fmt: skip
    holdings = write("ten.csv", [
        ["instrument_id", "tenedor", "valor_nominal_original_usd", "metodo", "fecha_dato",
         "source_id", "referencia"],
        ["INS-C-0001", "BCRA", 30e6, "directo", "2026-06-30", "SRC-0001", "x"],
        ["INS-C-0003", "BCRA", 5e6, "estimado", "2026-06-30", "SRC-0001", "x"],
    ])  # fmt: skip
    inst = write("inst.csv", [
        ["instrument_id", "instrument_family", "operation_name_sigade",
         "outstanding_usd_equiv_sigade"],
        ["INS-C-0001", "USD", "LETRA/U$S/BCRA", 30e6],
        ["INS-C-0002", "USD", "BONO GLOBAL/U$S/2030", 50e6],
        ["INS-C-0003", "USD", "BONO R.A./U$S/2030", 40e6],
        ["INS-C-0004", "multilateral", "BID", 20e6],
    ])  # fmt: skip
    head = ["instrument_id", "date", "base", "is_payment", "status", "currency_of_settlement",
            "monto_usd"]  # fmt: skip
    flows = write("flujos.csv", [head] + [
        [iid, "2026-07-09", "True", "True", "firme", "USD", 10e6]
        for iid in ("INS-C-0001", "INS-C-0002", "INS-C-0003", "INS-C-0004")
    ])  # fmt: skip
    out = tmp_path / "out"
    t = build_consolidation(comp, holdings, inst, flows, terms, date(2026, 6, 30), out)
    # Public USD 60: identified 30 (direct) + 5 (estimated); foreign law 10 goes to the
    # Global (10/50 = 20%); the remaining 15 to the free part of the local bond (15/35).
    assert t["bruta_musd"] == 40.0
    assert t["sector_publico_directo_musd"] == 10.0
    assert t["sector_publico_estimado_musd"] == 1.2  # 10 x 5/40, rounded
    assert t["proporcion_ley_extranjera_pct"] == 20.0
    assert t["proporcion_ley_nacional_pct"] == 42.86
    assert t["consolidada_musd"] == 23.0  # 40 - 10 - 1.25 - 2 - 3.75
    with (out / "necesidad_divisas_consolidada.csv").open(encoding="utf-8") as fh:
        rows = {r["instrument_id"]: r for r in csv.DictReader(fh)}
    assert float(rows["INS-C-0004"]["consolidada_musd"]) == 10.0
    assert float(rows["INS-C-0002"]["sector_publico_proporcional_musd"]) == 2.0
    assert float(rows["INS-C-0003"]["sector_publico_proporcional_musd"]) == 3.75
