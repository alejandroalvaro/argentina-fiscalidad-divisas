"""Verify Grok's SIGADE CSV exports against the original Access database stored in R2.

Usage:
    uv run python scripts/verify_sigade.py <sigade_zip_sha256> <grok_tablas_dir> [--xlsx <sha256>]

Requires mdbtools (`brew install mdbtools`). Checks:
1. Each table exported with mdb-export is byte-identical to Grok's CSV.
2. The sum of "Saldo en Dólares" (USD units) against cuadro A.1.1 (thousands of USD), if --xlsx.
3. Balances per operation against the sum of future principal in the maturity table.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import subprocess
import sys
import tempfile
import zipfile
from collections import defaultdict
from pathlib import Path

from argfx.config import get_settings
from argfx.storage import fetch_verified, get_store

TABLES = {
    "Saldo de Deuda en Situación de Pago Normal": "saldo_de_deuda_en_situacion_de_pago_normal",
    "Vencimientos de la Deuda en Situación de Pago Normal": (
        "vencimientos_de_la_deuda_en_situacion_de_pago_normal"
    ),
    "Deuda elegible pendiente de reestructuración": "deuda_elegible_pendiente_de_reestructuracion",
    "Deuda a Reestructurar": "deuda_a_reestructurar",
    "Tipos de Cambio": "tipos_de_cambio",
}


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("zip_sha256")
    ap.add_argument("grok_dir", type=Path)
    ap.add_argument(
        "--xlsx", help="sha256 of the Secretaría de Finanzas workbook (cuadros A.1-A.4)"
    )
    args = ap.parse_args()

    settings = get_settings()
    store = get_store(settings)
    ok = True
    with tempfile.TemporaryDirectory() as tmp:
        zpath = fetch_verified(store, args.zip_sha256, Path(tmp) / "sigade.zip", settings.cache_dir)
        with zipfile.ZipFile(zpath) as zf:
            (mdb_name,) = [n for n in zf.namelist() if n.lower().endswith(".mdb")]
            mdb = Path(zf.extract(mdb_name, tmp))

        exports: dict[str, bytes] = {}
        for table, slug in TABLES.items():
            ours = subprocess.run(["mdb-export", str(mdb), table], capture_output=True, check=True)
            exports[table] = ours.stdout
            (grok_file,) = args.grok_dir.glob(f"sigade_*_{slug}.csv")
            same = sha(ours.stdout) == sha(grok_file.read_bytes())
            ok &= same
            rows = ours.stdout.count(b"\n") - 1
            print(f"{'IDENTICAL' if same else 'DIFFERENT'}  {table} ({rows} rows)")

        def rows(table: str) -> list[dict[str, str]]:
            return list(csv.DictReader(exports[table].decode("utf-8").splitlines()))

        saldo = rows("Saldo de Deuda en Situación de Pago Normal")
        venc = rows("Vencimientos de la Deuda en Situación de Pago Normal")
        total_usd = sum(float(r["Saldo en Dólares"]) for r in saldo)
        print(f"Sum 'Saldo en Dólares' (USD): {total_usd:,.2f}")

        by_op: dict[str, float] = defaultdict(float)
        for r in saldo:
            by_op[r["Nombre de la Operación"]] += float(r["Saldo en Dólares"])
        for r in venc:
            by_op[r["Nombre de la Operación"]] -= float(r["Principal en Dólares"])
        gap = sum(by_op.values())
        worst = max(by_op.items(), key=lambda kv: abs(kv[1]))
        print(
            f"Balance minus future principal (USD): {gap:,.2f}; largest: {worst[1]:,.2f} {worst[0]}"
        )

        if args.xlsx:
            from openpyxl import load_workbook

            xpath = fetch_verified(store, args.xlsx, Path(tmp) / "deuda.xlsx", settings.cache_dir)
            ws = load_workbook(xpath, data_only=True, read_only=True)["A.1.1"]
            for row in ws.iter_rows(values_only=True):
                cells = [c for c in row if c is not None]
                if cells and "SITUACIÓN DE PAGO NORMAL" in str(cells[0]).upper():
                    a11 = float(cells[1])
                    match = abs(a11 * 1000 - total_usd) < 1
                    ok &= match
                    print(
                        f"A.1.1 (thousand USD): {a11:,.3f} -> x1000 "
                        f"{'matches' if match else 'DIFFERS'}"
                    )
                    break
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
