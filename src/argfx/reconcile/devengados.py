"""Interest capitalized and not paid, quarter by quarter, against the cash-basis fiscal result
(D-009).

The quarterly debt workbook of the Secretaría de Finanzas (A.2.2) has a line "Capitalización de
Bonos, Letras y otros préstamos", which, by its own footnote, includes the capitalization
accumulated on reopenings: when a LECAP or BONCAP is reopened, the new face value enters at
face and the capitalization it carries since issue enters through this line. That part was
paid in cash by the buyer; it is not interest accrued in the quarter. So:

    capitalized interest accrued = capitalization line - capitalization carried by reopenings

The second term comes from the placements of the Oficina Nacional de Crédito Público (sheets
of `colocaciones_<date>.xlsx`): for each placement of a fixed-rate LECAP or BONCAP settled
after its issue date, face x [(1 + TEM)^(DÍAS/30) - 1], with DÍAS from issue to settlement in
30E/360 (the legal formula). TAMAR and dual reopenings, whose TEM is not in the placement row,
use the effective value minus face as a proxy; rows without either are counted.

The fiscal side is the monthly cash-basis result of the Sector Público Nacional (primary
result, net interest, financial result), summed by quarter. The perimeters differ: A.2.2 is the
Administración Central with every holder, while net interest of the SPN excludes interest paid
inside the SPN (to the FGS, for instance). Both are reported as they are.

Units: thousands of millions of pesos (miles de millones), current pesos.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from argfx.contracts.calendar import _write
from argfx.contracts.daycount import year_fraction
from argfx.reconcile.finanzas import lines

TEM_IN_COUPON = re.compile(r"(\d+(?:,\d+)?)\s*%")


def quarter(d: date) -> str:
    return f"{d.year}T{(d.month - 1) // 3 + 1}"


def capitalization_line(xlsx: Path) -> tuple[str, float]:
    """Label and amount (miles de millones de pesos) of the A.2.2 capitalization line."""
    for line in lines(xlsx, "A.2.2"):
        if line.label.lower().startswith("capitaliz"):
            return line.label, line.values[1] / 1e6
    raise KeyError(f"{xlsx}: no capitalization line in A.2.2")


@dataclass
class Reopenings:
    fixed: float = 0.0  # computed with the TEM of the placement
    proxy: float = 0.0  # effective value minus face (TAMAR, duals)
    rows_fixed: int = 0
    rows_proxy: int = 0
    rows_missing: int = 0


def _as_date(v: object) -> date | None:
    return v.date() if isinstance(v, datetime) else None


def reopening_capitalization(workbooks: list[Path]) -> dict[str, Reopenings]:
    out: dict[str, Reopenings] = defaultdict(Reopenings)
    for path in workbooks:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb.sheetnames:
                header: list[str] | None = None
                for row in wb[sheet].iter_rows(values_only=True):
                    cells = list(row)
                    if header is None:
                        if any(isinstance(c, str) and c.startswith("Nombre") for c in cells):
                            header = [str(c).split("\n")[0].strip() if c else "" for c in cells]
                        continue
                    _placement(dict(zip(header, cells, strict=False)), out)
        finally:
            wb.close()
    return out


def _field(rec: dict[str, object], prefix: str) -> object:
    return next((v for k, v in rec.items() if k.startswith(prefix)), None)


def _placement(rec: dict[str, object], out: dict[str, Reopenings]) -> None:
    name = str(rec.get("Nombre del Instrumento") or "").upper()
    if "LECAP" not in name and "BONCAP" not in name:
        return
    issue = _as_date(rec.get("Fecha de emisión"))
    settle = _as_date(_field(rec, "Fecha coloc"))
    face = rec.get("Valor Nominal")
    if issue is None or settle is None or not isinstance(face, int | float) or settle <= issue:
        return  # a new issue carries no capitalization
    r = out[quarter(settle)]
    match = TEM_IN_COUPON.search(str(_field(rec, "Cupón") or ""))
    if match and "TAMAR" not in name and "DUAL" not in name:
        tem = float(match.group(1).replace(",", ".")) / 100
        factor = (1 + tem) ** (year_fraction(issue, settle, "30E/360") * 12)
        r.fixed += face * (factor - 1) / 1e3
        r.rows_fixed += 1
        return
    value = _field(rec, "Valor Efectivo")
    if not isinstance(value, int | float):
        value = _field(rec, "Valor Técnico")
    if isinstance(value, int | float):
        r.proxy += max(0.0, value - face) / 1e3
        r.rows_proxy += 1
    else:
        r.rows_missing += 1


def spn_quarterly(series_csv: Path) -> dict[str, dict[str, float]]:
    out: dict[str, dict[str, float]] = defaultdict(lambda: defaultdict(float))
    months: dict[str, set[str]] = defaultdict(set)
    with series_csv.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            d = date.fromisoformat(r["fecha"])
            q = quarter(d)
            out[q][r["series_id"]] += float(r["valor"]) / 1e3
            months[q].add(r["fecha"][:7])
    for q, ms in months.items():
        out[q]["meses"] = len(ms)
    return out


def build_accrual_table(
    quarterly: list[tuple[date, Path]],
    placements: list[Path],
    spn_csv: Path,
    out: Path,
) -> list[dict[str, object]]:
    reopenings = reopening_capitalization(placements)
    spn = spn_quarterly(spn_csv)
    rows = []
    for cutoff, xlsx in sorted(quarterly):
        q = quarter(cutoff)
        label, cap = capitalization_line(xlsx)
        reo = reopenings.get(q, Reopenings())
        accrued = cap - reo.fixed - reo.proxy
        fiscal = spn.get(q, {})
        complete = fiscal.get("meses") == 3
        financial = fiscal.get("spn_caja_resultado_financiero") if complete else None
        rows.append(
            {
                "trimestre": q,
                "capitalizacion_a22": round(cap, 1),
                "capitalizacion_en_ampliaciones_tasa_fija": round(reo.fixed, 1),
                "capitalizacion_en_ampliaciones_aproximada": round(reo.proxy, 1),
                "filas_sin_dato": reo.rows_missing,
                "interes_capitalizado_devengado": round(accrued, 1),
                "spn_resultado_primario_caja": _r(
                    fiscal.get("spn_caja_resultado_primario"), complete
                ),
                "spn_intereses_netos_caja": _r(fiscal.get("spn_caja_intereses_netos"), complete),
                "spn_resultado_financiero_caja": _r(financial, True),
                "resultado_financiero_con_capitalizado": (
                    None if financial is None else round(financial - accrued, 1)
                ),
                "rotulo_a22": label,
            }
        )
    _write(out / "devengados_trimestral.csv", rows)
    return rows


def _r(v: float | None, ok: bool) -> float | None:
    return round(v, 1) if ok and v is not None else None
