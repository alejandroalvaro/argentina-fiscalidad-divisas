"""Monthly exports by NCM position (ARCA, Información agregada de comercio exterior) and the duty
they owe at the rates of `alicuotas`.

Each monthly zip (lote AE, SRC-1375 onward) holds apostrophe-separated text files. The export
total by position is `total_expo_agregado_<YYYYMM>.lst`; older zips only have
`expo_agregado_<YYYY_MM>.lst`, by customs office and country, which is summed by position.
Only the export member is read, streaming: some zips also carry several GB of imports.

The estimated duty is FOB times the rate in force on the 15th of the month of shipment. It is an
accrual estimate by shipment; the law fixes the rate at the DJVE and the payment follows the
shipment by about 15 days (lote AD), so the estimate is compared with collections as a check,
not as a substitute.
"""

from __future__ import annotations

import io
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from argfx.simulation.alicuotas import RateBook

# Product groups for the outputs (first digits of the NCM position).
GROUPS = [
    ("soja", ("1201", "1208.10", "1507", "2304", "2302.50", "1517.90", "1518.00.90", "3826")),
    ("maiz", ("1005", "1102.20", "1103.13", "1104.23", "1108.12")),
    ("trigo y cebada", ("1001", "1003", "1101", "1107", "1109", "2302.30")),
    ("girasol", ("1206", "1512", "2306.30")),
    ("sorgo", ("1007",)),
    ("carnes", ("01", "02", "1601", "1602")),
    ("lacteos", ("04",)),
    ("pesca", ("03", "1604", "1605")),
    ("petroleo y combustibles", ("2709", "2710", "2711", "2707")),
    ("mineria", ("25", "26", "2836.91", "7106", "7108")),
]


def group_of(code: str) -> str:
    best = ("otros", 0)
    for name, prefixes in GROUPS:
        for p in prefixes:
            if code.startswith(p) and len(p) > best[1]:
                best = (name, len(p))
    return best[0]


def _member(zf: zipfile.ZipFile) -> str:
    names = zf.namelist()
    for prefix in ("total_expo", "expo"):
        # Some months (2017-12) nest the files in a YYYYMM/ folder.
        hits = [n for n in names if n.rsplit("/", 1)[-1].lower().startswith(prefix)]
        if hits:
            return hits[0]
    raise KeyError(f"no export file in {names}")


def read_month(path: Path) -> tuple[str, dict[str, float]]:
    """(YYYY-MM, FOB in USD by NCM position) from one monthly zip."""
    fob: dict[str, float] = defaultdict(float)
    month = ""
    with zipfile.ZipFile(path) as zf, zf.open(_member(zf)) as raw:
        header: list[str] | None = None
        for line in io.TextIOWrapper(raw, encoding="latin-1"):
            cells = [c.strip() for c in line.split("'")]
            if header is None:
                if "POS_NCM" in cells:
                    header = cells
                continue
            if len(cells) < len(header) or cells[0].startswith("-"):
                continue
            rec = dict(zip(header, cells, strict=False))
            try:
                fob[rec["POS_NCM"]] += float(rec["MONTO_FOB_DOLAR"])
            except ValueError:
                continue
            month = month or f"{rec['FECHA_'][:4]}-{rec['FECHA_'][4:6]}"
    return month, dict(fob)


@dataclass
class MonthDuty:
    month: str
    fob_usd: float = 0.0
    fob_with_rate_usd: float = 0.0
    duty_usd: float = 0.0
    by_group: dict[str, list[float]] | None = None  # group -> [fob, duty]


def estimate(month: str, fob: dict[str, float], book: RateBook) -> MonthDuty:
    day = date(int(month[:4]), int(month[5:7]), 15)
    out = MonthDuty(month, by_group=defaultdict(lambda: [0.0, 0.0]))
    for code, v in fob.items():
        out.fob_usd += v
        rate = book.rate(code[:10], day)
        cell = out.by_group[group_of(code)]  # type: ignore[index]
        cell[0] += v
        if rate is not None:
            out.fob_with_rate_usd += v
            out.duty_usd += v * rate[0] / 100
            cell[1] += v * rate[0] / 100
    return out


def crude_prices(xlsx: Path) -> dict[str, float]:
    """International price of crude (Decreto 488/2020 art. 7) in force each month (SRC-1566)."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        out: dict[str, float] = {}
        for row in wb.worksheets[0].iter_rows(values_only=True):
            cells = [c for c in row if c is not None]
            if len(cells) >= 3 and hasattr(cells[2], "year") and isinstance(cells[1], int | float):
                out.setdefault(f"{cells[2].year}-{cells[2].month:02d}", float(cells[1]))
        return out
    finally:
        wb.close()


def monthly_average(csv_path: Path) -> dict[str, float]:
    import csv

    acc: dict[str, list[float]] = defaultdict(list)
    with csv_path.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            acc[r["fecha"][:7]].append(float(r["valor"]))
    return {m: sum(v) / len(v) for m, v in acc.items()}


def observed_duties_usd(series_csv: Path, a3500: dict[str, float],
                        series_id: str = "142.3_DEREC_2001_M_20") -> dict[str, float]:  # fmt: skip
    """Collected export duties (millions of pesos, Recursos tributarios, lote D) in millions of
    USD at the monthly average of the A 3500."""
    import csv

    out = {}
    with series_csv.open(newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            m = r["fecha"][:7]
            if r["series_id"] == series_id and m in a3500:
                out[m] = float(r["valor"]) / a3500[m]
    return out


def build_duty_table(
    months: list[tuple[str, dict[str, float]]],
    book: RateBook,
    pi: dict[str, float],
    observed: dict[str, float],
    out: Path,
) -> list[dict[str, object]]:
    from argfx.contracts.calendar import _write
    from argfx.simulation.alicuotas import (
        FORMULA_488,
        FORMULA_488_FROM,
        crude_rate,
        formula_top,
    )

    rows, by_group = [], []
    for month, fob in sorted(months):
        e = estimate(month, {c: v for c, v in fob.items() if not c.startswith(FORMULA_488)},
                     book)  # fmt: skip
        formula = {c: v for c, v in fob.items() if c.startswith(FORMULA_488)}
        formula_fob = sum(formula.values())
        priced = month in pi and month >= FORMULA_488_FROM
        formula_duty = sum(
            v * crude_rate(pi[month], top=formula_top(c, month)) / 100 for c, v in formula.items()
        ) if priced else 0.0  # fmt: skip
        total = (e.duty_usd + formula_duty) / 1e6
        obs = observed.get(month)
        rows.append({
            "mes": month,
            "fob_musd": round((e.fob_usd + formula_fob) / 1e6, 1),
            "fob_con_alicuota_musd": round(
                (e.fob_with_rate_usd + (formula_fob if priced else 0.0)) / 1e6, 1),
            "derecho_estimado_musd": round(total, 1),
            "derecho_cobrado_musd": None if obs is None else round(obs, 1),
            "cobrado_sobre_estimado": None if obs is None or not total else round(obs / total, 3),
        })  # fmt: skip
        groups = dict(e.by_group or {})
        groups["petroleo y combustibles"] = [
            groups.get("petroleo y combustibles", [0.0, 0.0])[0] + formula_fob,
            groups.get("petroleo y combustibles", [0.0, 0.0])[1] + formula_duty,
        ]
        for g, (f, d) in sorted(groups.items()):
            by_group.append({"mes": month, "grupo": g, "fob_musd": round(f / 1e6, 2),
                             "derecho_estimado_musd": round(d / 1e6, 2)})  # fmt: skip
    _write(out / "derechos_estimados.csv", rows)
    _write(out / "derechos_estimados_por_grupo.csv", by_group)
    return rows
