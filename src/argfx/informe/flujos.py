"""Where pesos and dollars came in and went out over the mandate (D-016, D-017).

Two flow tables, one per currency, never added to each other:
- dollars: the BCRA balance cambiario (lote AH, SRC-1604), gross inflows and outflows of the
  foreign-exchange market, closed by the change in reserves from transactions. Millions of USD.
- pesos: the cash-basis result of the national public sector and its financial sources and uses
  (Secretaría de Hacienda, lote AM; financing table corrected in lote AN), closed by the
  financial result. Millions of pesos at the prices of the last month with a CPI index (INDEC,
  lote AM), deflated month by month.

The peso table keeps only pesos (D-017). Foreign-currency debt, its interest (Secretaría de
Finanzas boletín, sheet A.5, by currency of payment) and the Treasury's dollar deposits at the
BCRA leave it; what is left of them on the peso side is a single crossing, the pesos applied
to dollar obligations, net. A third table shows the Treasury's dollars, in millions of USD.

Each table is a list of (origen, destino, monto) links around a single hub node, ready for a
Sankey diagram. Inflows go origin -> hub and outflows hub -> destination; both sides add up to
the same total, and any gap is reported, never absorbed.
"""

from __future__ import annotations

import csv
import datetime as dt
import json
from dataclasses import dataclass
from pathlib import Path

HUB_USD = "Mercado de cambios"
HUB_ARS = "Sector Público Nacional"
HUB_TESORO_USD = "Tesoro Nacional (dólares)"
# Debt written in dollars and paid in pesos (BONAR DUAL, BONTE and LETES "paga en pesos",
# LELINK) is a peso flow (D-017). The AIF books it as foreign-currency debt, by the currency
# it is written in: cash placed at auctions explains 96% to 101% of the AIF foreign-currency
# borrowing in 2025-09 and 2025-10 (lote AQ). Its cash amortizations and placements (lote AR)
# leave the dollar tables and enter the peso table as their own lines.
# The crossing between the two tables: pesos the Treasury applied to dollar obligations, net,
# in three parts. Dollars bought from the BCRA or in the market (lote AI); dollar debt held by
# the BCRA cancelled with pesos (Letras Intransferibles, RC 27/2026); and the rest, which the
# sources do not open: timing differences, gross renewals and the months without lote AP.
CROSS_ARS = {
    "compras": "Compra de dólares (al BCRA y en el mercado)",
    "cancelado": "Letras Intransferibles canceladas con pesos",
    "resto": "Diferencia entre monedas que la fuente no abre (neto)",
}
CROSS_USD = {
    "compras": "Comprados con pesos",
    "cancelado": "Letras Intransferibles canceladas con pesos",
    "resto": "Diferencia entre monedas que la fuente no abre (neto)",
}
# Letras Intransferibles bought back with pesos, millions of pesos (RC 27/2026, SRC-2031: $18,4
# billones of the $24,4 billones of BCRA dividends, SRC-2047). The buybacks of 2025-05 and
# 2025-08 used IMF dollars, and the one of 2026-08 dollars from the BCRA for BOPREAL (lote AO).
CANCELLED_WITH_PESOS = {"2026-05": 18_400_000.0}
PAID_IN_PESOS = "Deuda escrita en dólares, pagada en pesos"
PLACED_FOR_PESOS = "Deuda escrita en dólares, colocada por pesos"
GAP = "Sin identificar en la fuente"

# Balance cambiario columns (sheet "Balance Cambiario", codes on row 11).
USD_IN = {
    "Exportaciones de bienes": "bal003",
    "Servicios (cobros)": "bal006",
    "Intereses cobrados": "bal010",
    "Utilidades cobradas": "bal017",
    "Ingreso secundario (cobros)": "bal020",
    "Inversión directa (ingresos)": "bal026",
    "Portafolio (ingresos)": "bal029",
    "Préstamos y títulos (ingresos)": "bal032",
    "FMI (desembolsos)": "bal035",
    "Otros organismos (desembolsos)": "bal038",
    "Venta de billetes y divisas de privados": "bal041",
}
USD_OUT = {
    "Importaciones de bienes": "bal004",
    "Servicios (pagos)": "bal007",
    "Intereses al FMI": "bal012",
    "Intereses a otros organismos": "bal013",
    "Otros intereses": "bal014",
    "Intereses del Gobierno Nacional": "bal015",
    "Utilidades y dividendos": "bal018",
    "Ingreso secundario (pagos)": "bal021",
    "Inversión directa (egresos)": "bal027",
    "Portafolio (egresos)": "bal030",
    "Préstamos y títulos (egresos)": "bal033",
    "FMI (capital)": "bal036",
    "Otros organismos (capital)": "bal039",
    "Compra de billetes y divisas de privados": "bal042",
}
# Net items: an inflow when positive, an outflow when negative.
USD_NET = {
    "Cuenta capital": "bal023",
    "Canje": "bal043",
    "Activos externos del sector financiero": "bal044",
    "Activos externos del sector público": "bal045",
    "Compra-venta de títulos": "bal046",
    "Otras operaciones del sector público": "bal047",
    "Otros movimientos": "bal048",
    "Concepto no informado": "bal049",
}
USD_RESERVES = "bal050"


@dataclass
class Flows:
    moneda: str
    unidad: str
    desde: str
    hasta: str
    hub: str
    entradas: dict[str, float]
    salidas: dict[str, float]
    clave: str = ""  # file name; defaults to the currency

    @property
    def descuadre(self) -> float:
        return sum(self.entradas.values()) - sum(self.salidas.values())

    def links(self) -> list[tuple[str, str, float]]:
        """Sankey links; a gap between the two sides is drawn as its own node."""
        out = [(k, self.hub, v) for k, v in self.entradas.items() if v > 0]
        out += [(self.hub, k, v) for k, v in self.salidas.items() if v > 0]
        gap = self.descuadre
        tol = 1e-6 * sum(self.entradas.values())  # rounding in the source tables
        if gap < -tol:
            out.append((GAP, self.hub, -gap))
        elif gap > tol:
            out.append((self.hub, GAP, gap))
        return out


def _add_signed(name: str, value: float, ins: dict[str, float], outs: dict[str, float]) -> None:
    if value >= 0:
        ins[name] = ins.get(name, 0.0) + value
    else:
        outs[name] = outs.get(name, 0.0) - value


def usd_flows(monthly: dict[str, dict[str, float]], desde: str, hasta: str) -> Flows:
    """`monthly` maps YYYY-MM to balance cambiario codes (millions of USD)."""
    months = sorted(m for m in monthly if desde <= m <= hasta)
    if not months:
        raise ValueError(f"sin datos del balance cambiario entre {desde} y {hasta}")

    def total(code: str) -> float:
        return sum(monthly[m].get(code) or 0.0 for m in months)

    ins = {k: total(c) for k, c in USD_IN.items()}
    outs = {k: abs(total(c)) for k, c in USD_OUT.items()}
    for k, c in USD_NET.items():
        _add_signed(k, total(c), ins, outs)
    reserves = total(USD_RESERVES)
    if reserves >= 0:
        outs["Aumento de reservas"] = reserves
    else:
        ins["Caída de reservas"] = -reserves
    return Flows("USD", "millones de USD", months[0], months[-1], HUB_USD, ins, outs)


def read_balance_cambiario(xlsx: Path) -> dict[str, dict[str, float]]:
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["Balance Cambiario"].iter_rows(values_only=True))
    codes = {j: c for j, c in enumerate(rows[10]) if isinstance(c, str) and c.startswith("bal")}
    out: dict[str, dict[str, float]] = {}
    for r in rows[11:]:
        if isinstance(r[0], dt.datetime):
            out[f"{r[0].year}-{r[0].month:02d}"] = {
                c: float(r[j]) for j, c in codes.items() if isinstance(r[j], (int, float))
            }
    return out


# AIF columns (lote AM), millions of current pesos. Revenue and primary spending come from the
# result table; financing from the sources and uses table, total national public sector, so that
# both sides cover the same perimeter.
SPENDING = {
    "Prestaciones sociales": "prestaciones_sociales",
    "Subsidios económicos": "subsidios_economicos",
    "Transferencias a provincias": "transferencias_corrientes_a_provincias",
    "Universidades": "transferencias_a_universidades",
}
FX_DISBURSED = "endeudamiento_moneda_extranjera_total_spn"
FX_AMORTIZED = "amortizacion_moneda_extranjera_total_spn"
FIN_IN = {
    "Deuda en pesos colocada": "endeudamiento_moneda_local_total_spn",
    "Uso de activos financieros": "disminucion_inversion_financiera_total_spn",
    "Otros pasivos (aumento)": "incremento_otros_pasivos_total_spn",
    # Excluded from revenue by note (1) of each workbook, so only a financial source (lote AN).
    "Utilidades del BCRA (rentas percibidas)": "rentas_percibidas_del_bcra_total_spn",
}
FIN_OUT = {
    "Deuda en pesos amortizada": "amortizacion_moneda_local_total_spn",
    "Activos financieros (aumento)": "inversion_financiera_total_spn",
    "Otros pasivos (disminución)": "disminucion_otros_pasivos_total_spn",
}
SUFFIX = "_millones_pesos_corrientes"


def _num(row: dict[str, str], col: str) -> float:
    v = row.get(col + SUFFIX, row.get(col, ""))
    return float(v) if v not in ("", None) else 0.0


@dataclass
class FxMonth:
    """The dollar side of one month: average A 3500 rate (pesos per USD), foreign-currency
    interest paid (millions of USD) and the change in the Treasury's dollar deposits at the
    BCRA (millions of USD, positive when they grow)."""

    tipo_de_cambio: float
    intereses_usd: float
    depositos_usd: float
    compras_usd: float = 0.0  # dollars the Treasury bought with pesos
    cancelado_pesos: float = 0.0  # dollar debt cancelled with pesos, millions of pesos
    # Debt written in dollars and paid in pesos, millions of USD (lote AR): amortized in cash
    # and placed for cash at auctions. The AIF books both as foreign-currency debt.
    pagado_en_pesos_usd: float = 0.0
    colocado_por_pesos_usd: float = 0.0


def ars_month(
    r: dict[str, str], f: dict[str, str], fx: FxMonth
) -> tuple[dict[str, float], dict[str, float]]:
    """One month of the peso circuit, in current pesos: (inflows, outflows)."""
    tax, dex, rent = (
        _num(r, c)
        for c in ("ingresos_tributarios", "derechos_de_exportacion", "rentas_de_la_propiedad")
    )
    ins = {
        "Derechos de exportación": dex,
        "Otros impuestos y aportes": tax - dex,
        "Rentas de la propiedad": rent,
        "Otros ingresos corrientes y de capital": _num(r, "ingresos_totales") - tax - rent,
    }
    outs = {k: _num(r, c) for k, c in SPENDING.items()}
    outs["Funcionamiento y otros gastos corrientes"] = _num(r, "gastos_corrientes_primarios") - sum(
        outs.values()
    )
    outs["Gasto de capital"] = _num(r, "gastos_de_capital")
    ins.update({k: _num(f, c) for k, c in FIN_IN.items()})
    outs.update({k: _num(f, c) for k, c in FIN_OUT.items()})
    fx_interest = fx.intereses_usd * fx.tipo_de_cambio
    deposits = fx.depositos_usd * fx.tipo_de_cambio
    outs["Intereses de la deuda en pesos (netos)"] = _num(r, "intereses_netos") - fx_interest
    # The dollar deposits sit inside the financial assets of the sources and uses table.
    if deposits >= 0:
        outs["Activos financieros (aumento)"] -= deposits
    else:
        ins["Uso de activos financieros"] += deposits
    paid = fx.pagado_en_pesos_usd * fx.tipo_de_cambio
    placed = fx.colocado_por_pesos_usd * fx.tipo_de_cambio
    outs[PAID_IN_PESOS] = paid
    ins[PLACED_FOR_PESOS] = placed
    amortized, disbursed = _num(f, FX_AMORTIZED) - paid, _num(f, FX_DISBURSED) - placed
    cross = amortized + fx_interest + deposits - disbursed
    bought = fx.compras_usd * fx.tipo_de_cambio
    outs[CROSS_ARS["compras"]] = bought
    outs[CROSS_ARS["cancelado"]] = fx.cancelado_pesos
    outs[CROSS_ARS["resto"]] = cross - bought - fx.cancelado_pesos
    return ins, outs


def _months(desde: str, hasta: str, *tables: dict[str, object]) -> list[str]:
    first = tables[0]
    return sorted(m for m in first if desde <= m <= hasta and all(m in t for t in tables[1:]))


def _split(net: dict[str, float]) -> tuple[dict[str, float], dict[str, float]]:
    """Signed totals over the window (inflow positive) into the two sides of the hub."""
    ins = {n: v for n, v in net.items() if v > 0}
    outs = {n: -v for n, v in net.items() if v < 0}
    return ins, outs


def ars_flows(
    results: list[dict[str, str]],
    financing: list[dict[str, str]],
    cpi: dict[str, float],
    fx: dict[str, FxMonth],
    desde: str,
    hasta: str,
) -> Flows:
    """Pesos of the national public sector, at prices of the last month of the window.

    Only months present in every input are used; the window shrinks to them. Each line is
    netted over the window, so a line sits on one side of the hub only.
    """
    res = {r["mes"]: r for r in results}
    fin = {r["mes"]: r for r in financing}
    months = _months(desde, hasta, res, fin, cpi, fx)
    if not months:
        raise ValueError(
            f"sin meses con resultado, financiamiento, IPC y divisas entre {desde} y {hasta}"
        )
    base = cpi[months[-1]]
    net: dict[str, float] = {}
    for m in months:
        k = base / cpi[m]
        m_ins, m_outs = ars_month(res[m], fin[m], fx[m])
        for name, v in m_ins.items():
            net[name] = net.get(name, 0.0) + k * v
        for name, v in m_outs.items():
            net[name] = net.get(name, 0.0) - k * v
    ins, outs = _split(net)
    return Flows(
        "ARS", f"millones de pesos de {months[-1]}", months[0], months[-1], HUB_ARS, ins, outs
    )


def treasury_usd_flows(
    financing: list[dict[str, str]], fx: dict[str, FxMonth], desde: str, hasta: str
) -> Flows:
    """The Treasury's dollars, millions of USD: debt in foreign currency, its interest, the
    deposits at the BCRA and what pesos covered. Disbursements and amortizations,
    recorded in pesos by the sources and uses table, are converted at the month's average rate."""
    fin = {r["mes"]: r for r in financing}
    months = _months(desde, hasta, fin, fx)
    if not months:
        raise ValueError(f"sin meses con financiamiento y divisas entre {desde} y {hasta}")
    net: dict[str, float] = {}

    def add(name: str, v: float) -> None:
        net[name] = net.get(name, 0.0) + v

    for m in months:
        x, rate = fx[m], fx[m].tipo_de_cambio
        disbursed = _num(fin[m], FX_DISBURSED) / rate - x.colocado_por_pesos_usd
        amortized = _num(fin[m], FX_AMORTIZED) / rate - x.pagado_en_pesos_usd
        add("Deuda en divisas desembolsada", disbursed)
        add("Deuda en divisas amortizada", -amortized)
        add("Intereses en divisas", -x.intereses_usd)
        add("Depósitos en dólares en el BCRA (variación)", -x.depositos_usd)
        cross = amortized + x.intereses_usd + x.depositos_usd - disbursed
        cancelled = x.cancelado_pesos / rate
        add(CROSS_USD["compras"], x.compras_usd)
        add(CROSS_USD["cancelado"], cancelled)
        add(CROSS_USD["resto"], cross - x.compras_usd - cancelled)
    ins, outs = _split(net)
    return Flows(
        "USD", "millones de USD", months[0], months[-1], HUB_TESORO_USD, ins, outs, "usd_tesoro"
    )


def read_a3500(path: Path) -> tuple[dict[str, float], dict[str, float]]:
    """(monthly average, last observation of the month) of the A 3500 rate."""
    acc: dict[str, list[tuple[str, float]]] = {}
    for r in read_csv(path):
        acc.setdefault(r["fecha"][:7], []).append((r["fecha"], float(r["valor"])))
    avg = {m: sum(v for _, v in xs) / len(xs) for m, xs in acc.items()}
    eom = {m: max(xs)[1] for m, xs in acc.items()}
    return avg, eom


def read_fx_interest(xlsx: Path) -> dict[str, float]:
    """Interest paid in foreign currency, millions of USD by month (boletín, sheet A.5, rows
    "MONEDA EXTRANJERA" then "INTERES")."""
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    rows = list(wb["A.5"].iter_rows(values_only=True))
    header = next(r for r in rows if sum(isinstance(c, dt.datetime) for c in r) > 6)
    cols = {
        j: f"{c.year}-{c.month:02d}" for j, c in enumerate(header) if isinstance(c, dt.datetime)
    }
    labels = [next((str(c).strip() for c in r if c is not None), "") for r in rows]
    start = labels.index("- MONEDA EXTRANJERA")
    i = next(i for i in range(start + 1, len(rows)) if labels[i] == "INTERES")
    return {m: float(rows[i][j]) for j, m in cols.items() if isinstance(rows[i][j], (int, float))}


def fx_months(
    a3500: tuple[dict[str, float], dict[str, float]],
    fx_interest: dict[str, float],
    deposits: list[dict[str, str]],
    purchases: dict[str, float] | None = None,
    paid_in_pesos: dict[str, tuple[float, float]] | None = None,
) -> dict[str, FxMonth]:
    """Join rate, interest, deposits and the Treasury's purchases of dollars. Deposits come in
    pesos at the end of each month; they are turned into dollars at that day's rate before
    taking the monthly change."""
    avg, eom = a3500
    col = "gobierno_nacional_moneda_extranjera_expresada_en_pesos_millones"
    usd = {
        r["mes"]: float(r[col]) / eom[r["mes"]] for r in deposits if r.get(col) and r["mes"] in eom
    }
    out = {}
    for m in sorted(usd):
        prev = _previous(m)
        if prev in usd and m in avg and m in fx_interest:
            out[m] = FxMonth(
                avg[m],
                fx_interest[m],
                usd[m] - usd[prev],
                (purchases or {}).get(m, 0.0),
                CANCELLED_WITH_PESOS.get(m, 0.0),
                *(paid_in_pesos or {}).get(m, (0.0, 0.0)),
            )
    return out


def read_purchases(path: Path) -> dict[str, float]:
    """Dollars the Treasury bought, millions of USD by month (lote AI): directly from the BCRA
    plus net in the market, as the BCRA reports them. Months marked n/d count as unknown and are
    left out, so the rest of the crossing absorbs them."""
    out = {}
    for r in read_csv(path):
        try:
            out[r["mes"]] = float(r["compras_directas_al_bcra_usd_millones_informe_bcra"]) + float(
                r["compras_netas_tesoro_en_mercado_usd_millones_informe_bcra"]
            )
        except ValueError:
            continue
    return out


def read_paid_in_pesos(amortized: Path, placed: Path) -> dict[str, tuple[float, float]]:
    """(amortized in cash, placed for cash) of debt written in dollars and paid in pesos,
    millions of USD by month (lote AR). Amortizations: A.4 of the Finanzas boletín less the
    swaps without cash of the ONCP sheet, counting as cash the early redemption of 2023-12,
    which the AIF books as cash. Placements: cash at auctions (ONCP), with the BONTE DUAL of
    2026-07 (USD 3.148 million): leaving it out leaves the same amount of dollars without use in
    that month, so the AIF books it as foreign-currency debt too."""
    out: dict[str, list[float]] = {}
    for r in read_csv(amortized):
        out.setdefault(r["mes"], [0.0, 0.0])[0] = float(
            r["caja_estimada_maxima_incluye_inseparable"] or 0
        )
    for r in read_csv(placed):
        out.setdefault(r["mes"], [0.0, 0.0])[1] = float(r["total_con_bonte_dual_usd_millones"] or 0)
    return {m: (a, p) for m, (a, p) in out.items()}


def _previous(m: str) -> str:
    y, mo = int(m[:4]), int(m[5:])
    return f"{y - 1}-12" if mo == 1 else f"{y}-{mo - 1:02d}"


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def read_cpi(path: Path) -> dict[str, float]:
    out = {}
    for r in read_csv(path):
        idx = next((r[c] for c in r if c.startswith("indice")), "")
        if idx not in ("", None):
            out[r["mes"]] = float(idx)
    return out


def write_flows(flows: Flows, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"flujos_{flows.clave or flows.moneda.lower()}_{flows.desde}_{flows.hasta}.csv"
    with path.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["origen", "destino", "monto", "unidad"])
        for a, b, v in flows.links():
            w.writerow([a, b, round(v, 1), flows.unidad])
    return path


PLOTLY = "https://cdnjs.cloudflare.com/ajax/libs/plotly.js/2.35.2/plotly.min.js"


def sankey_html(flows: Flows, title: str) -> str:
    links = flows.links()
    nodes = list(dict.fromkeys([a for a, _, _ in links] + [b for _, b, _ in links]))
    pos = {n: i for i, n in enumerate(nodes)}
    color = "#4E80EE" if flows.moneda == "USD" else "#E1B53E"
    data = {
        "type": "sankey",
        "valueformat": ",.0f",
        "valuesuffix": f" {flows.unidad}",
        "node": {"label": nodes, "pad": 14, "thickness": 16, "color": "#20293A"},
        "link": {
            "source": [pos[a] for a, _, _ in links],
            "target": [pos[b] for _, b, _ in links],
            "value": [round(v, 1) for _, _, v in links],
            "color": [color + "66"] * len(links),
        },
    }
    layout = {
        "title": {"text": title},
        "font": {"family": "DM Sans, system-ui, sans-serif", "size": 12},
        "height": 760,
    }
    return (
        "<!doctype html><html lang='es'><head><meta charset='utf-8'>"
        f"<title>{title}</title><script src='{PLOTLY}'></script></head><body>"
        "<div id='s'></div><script>"
        f"Plotly.newPlot('s',[{json.dumps(data, ensure_ascii=False)}],"
        f"{json.dumps(layout, ensure_ascii=False)},{{responsive:true}});"
        "</script></body></html>"
    )
